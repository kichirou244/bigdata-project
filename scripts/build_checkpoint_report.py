"""Sinh Checkpoint 2 report PPTX từ template của GV (không sửa template gốc)."""
from __future__ import annotations

import copy
import subprocess
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"D:\Edu\BD\BigData_Buoi06_DoAn2_Checkpoint2_Report_Template.pptx")
OUT = ROOT / "docs" / "Checkpoint2_Report_YouTube_Trending.pptx"

EMU_IN = 914400
BODY_FONT, MONO_FONT = "Calibri", "Consolas"

NAVY = RGBColor(0x0B, 0x1F, 0x33)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GRAY = RGBColor(0x5B, 0x6B, 0x7A)
MUTED = RGBColor(0x8A, 0x9A, 0xA8)
GREEN = RGBColor(0x0E, 0x7C, 0x4A)
LIGHT = RGBColor(0xD5, 0xE3, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xDB, 0xE3, 0xEA)
CK_FILL = RGBColor(0xEF, 0xF6, 0xFF)

TITLE = "Phân tích Xu hướng & Tốc độ Tăng trưởng Video Thịnh hành YouTube"

X_LEFT, X_RIGHT, CARD_W = 0.85, 6.85, 5.55
FULL_X, FULL_W = 0.85, 11.63
TOP, BOTTOM = 1.95, 6.95
PAD, TITLE_DY, TITLE_H, BODY_DY, PAD_B, GAP = 0.28, 0.14, 0.34, 0.60, 0.16, 0.20

FONT_FILES = {
    ("calibri", False): r"C:\Windows\Fonts\calibri.ttf",
    ("calibri", True): r"C:\Windows\Fonts\calibrib.ttf",
    ("consolas", False): r"C:\Windows\Fonts\consola.ttf",
    ("consolas", True): r"C:\Windows\Fonts\consolab.ttf",
}
PX_PER_PT, SAFETY = 8, 1.06
LINE_EM = {"calibri": 1.22, "consolas": 1.17}

try:
    from PIL import ImageFont

    HAVE_PIL = True
except ImportError:
    HAVE_PIL = False

_fonts: dict[tuple[str, bool, float], object] = {}


def _font(mono: bool, bold: bool, size: float):
    key = ("consolas" if mono else "calibri", bold, size)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(FONT_FILES[key[:2]], int(round(size * PX_PER_PT)))
    return _fonts[key]


def wrap_lines(text: str, width: float, size: float, mono: bool, bold: bool = False) -> int:
    if not text.strip():
        return 1
    if not HAVE_PIL:
        cpl = max(1, int(width * 72 / (size * (0.55 if mono else 0.50))))
        return -(-len(text) // cpl)
    font = _font(mono, bold, size)
    limit = width * 72 * PX_PER_PT
    lines, cur = 1, ""
    for word in text.split(" "):
        cand = f"{cur} {word}" if cur else word
        if font.getlength(cand) <= limit or not cur:
            cur = cand
        else:
            lines += 1
            cur = word
    return lines


def spec(line) -> tuple[str, bool]:
    if isinstance(line, tuple):
        return line[0], (line[2] if len(line) > 2 else False)
    return line, False


def measure(lines, width: float, size: float, mono: bool) -> float:
    n = sum(wrap_lines(spec(x)[0], width, size, mono, spec(x)[1]) for x in lines)
    return n * size * LINE_EM["consolas" if mono else "calibri"] / 72 * SAFETY


def emu(inches: float) -> Emu:
    return Emu(int(round(inches * EMU_IN)))


def by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(name)


def drop(slide, name: str) -> None:
    shape = by_name(slide, name)
    shape._element.getparent().remove(shape._element)


def place(shape, *, left=None, top=None, width=None, height=None) -> None:
    for attr, value in (("left", left), ("top", top), ("width", width), ("height", height)):
        if value is not None:
            setattr(shape, attr, emu(value))


def donor_rpr(paragraph):
    for run in paragraph.runs:
        rPr = run._r.find(qn("a:rPr"))
        if rPr is not None:
            return copy.deepcopy(rPr)
    return None


def set_lines(shape, lines, *, mono=False, size=12.0, color=GRAY) -> None:
    tf = shape.text_frame
    first = tf.paragraphs[0]
    rpr = donor_rpr(first)
    xml = copy.deepcopy(first._p)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)

    for i, item in enumerate(lines):
        text, bold = spec(item)
        col = item[1] if isinstance(item, tuple) else color
        if i == 0:
            p = first
        else:
            p_xml = copy.deepcopy(xml)
            for child in list(p_xml):
                if child.tag != qn("a:pPr"):
                    p_xml.remove(child)
            tf._txBody.append(p_xml)
            p = tf.paragraphs[-1]
        run = p.add_run()
        run.text = text
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        run.font.name = MONO_FONT if mono else BODY_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = col


def fit(lines, width: float, size: float, mono: bool, avail: float) -> float:
    while size > 8.5 and measure(lines, width, size, mono) > avail:
        size -= 0.25
    return size


def card(slide, n: int, title: str, lines, *, x=X_LEFT, w=CARD_W, top=TOP, h=1.2,
         size=11.0, mono=False, min_size=8.5, title_size=16.0) -> float:
    shape = by_name(slide, f"Shape {n}")
    place(shape, left=x, top=top, width=w, height=h)
    place(by_name(slide, f"Text {n + 1}"), left=x + PAD, top=top + TITLE_DY,
          width=w - 2 * PAD, height=TITLE_H)
    title_shape = by_name(slide, f"Text {n + 1}")
    set_lines(title_shape, [(title, BLUE, True)], size=title_size)
    title_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    body = by_name(slide, f"Text {n + 2}")
    body_w = w - 2 * PAD
    body_top = top + BODY_DY
    avail = max(0.3, top + h - PAD_B - body_top)
    used = fit(lines, body_w, size, mono, avail)
    place(body, left=x + PAD, width=body_w)
    set_lines(body, lines, mono=mono, size=used)
    body.text_frame.word_wrap = True
    body.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    content_h = measure(lines, body_w, used, mono)
    body.height = emu(min(content_h, avail))
    place(body, top=body_top)
    return used


def natural(c: dict, w: float, shrink: float = 0.0) -> float:
    size = max(8.5, c.get("size", 11.0) - shrink)
    h = measure(c["lines"], w - 2 * PAD, size, c.get("mono", False))
    return max(h + BODY_DY + PAD_B, 1.05)


def stack(slide, cards, *, x=X_LEFT, w=CARD_W, top=TOP, bottom=BOTTOM) -> None:
    avail = bottom - top - GAP * (len(cards) - 1)
    shrink = 0.0
    nat = [natural(c, w) for c in cards]
    while sum(nat) > avail and shrink < 4.0:
        shrink += 0.25
        nat = [natural(c, w, shrink) for c in cards]
    extra = max(0.0, (avail - sum(nat)) / len(cards))
    y = top
    for c, h in zip(cards, nat):
        card(slide, c["n"], c["title"], c["lines"],
             x=x, w=w, top=y, h=h + extra,
             size=max(8.5, c.get("size", 11.0) - shrink),
             mono=c.get("mono", False), min_size=c.get("min_size", 8.5))
        y += h + extra + GAP


def grid(slide, cells, *, bottom=BOTTOM) -> None:
    avail = bottom - TOP - GAP
    shrink = 0.0
    while shrink < 4.0:
        nat = [natural(c, CARD_W, shrink) for c in cells]
        if max(nat[0], nat[1]) + max(nat[2], nat[3]) <= avail:
            break
        shrink += 0.25
    nat = [natural(c, CARD_W, shrink) for c in cells]
    row1, row2 = max(nat[0], nat[1]), max(nat[2], nat[3])
    extra = max(0.0, (avail - row1 - row2) / 2)
    row1, row2 = row1 + extra, row2 + extra
    slots = [(cells[0], row1, TOP, X_LEFT), (cells[1], row1, TOP, X_RIGHT),
             (cells[2], row2, TOP + row1 + GAP, X_LEFT),
             (cells[3], row2, TOP + row1 + GAP, X_RIGHT)]
    for c, h, y, x in slots:
        card(slide, c["n"], c["title"], c["lines"], x=x, w=CARD_W, top=y, h=h,
             size=max(8.5, c.get("size", 11.0) - shrink), mono=c.get("mono", False))


def add_shape(slide, kind, x, y, w, h, fill_color, line_color=None, line_w=1.0, radius=None):
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w)
    if radius is not None:
        shape.adjustments[0] = radius
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return shape


def add_text(slide, x, y, w, h, lines, *, size=12.0, color=GRAY, mono=False,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    set_lines(box, lines, mono=mono, size=size, color=color)
    box.height = emu(measure(lines, w, size, mono))
    return box


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True,
                          check=True).stdout.strip()


def build() -> int:
    from pptx import Presentation

    if not TEMPLATE.exists():
        raise SystemExit(f"Không tìm thấy template: {TEMPLATE}")
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)
    tag = git("describe", "--tags", "--abbrev=0")

    s = slides[0]
    for name in ("Shape 5", "Text 6", "Text 7"):
        drop(s, name)
    place(by_name(s, "Text 0"), width=11.0)
    set_lines(by_name(s, "Text 0"), ["ĐỒ ÁN 2 • CHECKPOINT 2"], size=13, color=LIGHT)
    for name, text in (
        ("Text 2", "Nhóm: YouTube Trending Analytics"),
        ("Text 3", f"Đề tài: {TITLE}"),
        ("Text 4", f"Tag: {tag}"),
    ):
        shape = by_name(s, name)
        place(shape, width=11.0, height=0.30)
        set_lines(shape, [text], size=15.5, color=LIGHT)
    add_shape(s, MSO_SHAPE.RECTANGLE, 0.75, 6.34, 11.83, 0.014, RGBColor(0x1E, 0x3A, 0x52))
    add_text(
        s, 0.75, 6.48, 11.83, 0.3,
        ["Buổi 6 · 20/10/2026   |   Phan Đình Bảo · Trương Đình Phúc · "
         "Trần Thị Khánh Linh · Đoàn Ngọc Trân Châu"],
        size=12, color=RGBColor(0x93, 0xAD, 0xC2),
    )

    s = slides[1]
    for name in ("Text 2", "Text 5"):
        drop(s, name)
    place(by_name(s, "Shape 3"), left=FULL_X, top=TOP, width=FULL_W, height=4.2)
    place(by_name(s, "Text 4"), left=1.13, top=TOP + TITLE_DY, width=FULL_W - 0.56, height=TITLE_H)
    set_lines(by_name(s, "Text 4"), ["SƠ ĐỒ ARCHITECTURE — phần đã triển khai"], size=16,
              color=BLUE)

    stages = [
        ("01", "Kaggle CSV", ["8,59 GB · 113 nước", "6.094.047 record"]),
        ("02", "ingest_file.py", ["quality gate + SHA-256", "copy nguyên bản"]),
        ("03", "Raw storage", ["data/raw/…/run_…/", "+ manifest.json"]),
    ]
    box_w, gap, box_y, box_h = 3.39, 0.45, 2.78, 2.45
    for i, (num, name, details) in enumerate(stages):
        bx = 1.13 + i * (box_w + gap)
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, box_y, box_w, box_h,
                  CK_FILL, BLUE, 1.0, radius=0.06)
        badge = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx + 0.2, box_y + 0.24, 0.5, 0.42,
                          BLUE, None, radius=0.25)
        set_lines(badge, [(num, WHITE, True)], size=13)
        badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in badge.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
        add_text(s, bx + 0.2, box_y + 0.85, box_w - 0.4, 0.4, [(name, NAVY, True)], size=15)
        add_text(s, bx + 0.2, box_y + 1.35, box_w - 0.4, 0.7,
                 [(d, GRAY) for d in details], size=11)
        if i < len(stages) - 1:
            add_shape(s, MSO_SHAPE.RIGHT_ARROW, bx + box_w + 0.09, box_y + box_h / 2 - 0.12,
                      0.27, 0.24, MUTED)

    place(by_name(s, "Text 6"), left=1.13, top=5.55, width=FULL_W - 0.56, height=0.5)
    set_lines(by_name(s, "Text 6"), [
        "Trade-off: copy nguyên bản (tốn 8,59 GB đĩa) để đổi lấy khả năng truy vết "
        "và xử lý lại từ raw."
    ], size=11, color=GRAY)

    s = slides[2]
    drop(s, "Text 2")
    stack(s, [dict(n=3, title="EVIDENCE", size=12.0, mono=True, lines=[
        "$ python src\\env_check.py",
        "",
        "Python   : 3.12.10",
        "Java     : 17.0.20.1",
        "Spark    : 3.5.1",
        "Count    : 1000000",
        ("RESULT   : PASS (exit 0)", GREEN, True),
    ])], x=X_LEFT, w=CARD_W)
    stack(s, [
        dict(n=6, title="SETUP COMMAND", size=10.5, mono=True, lines=[
            "python -m venv .venv",
            ".venv\\Scripts\\pip install -r requirements.txt",
            ".venv\\Scripts\\python.exe src\\env_check.py",
        ]),
        dict(n=9, title="EXECUTION MODE", size=11.5, lines=[
            "local[*] · Windows · Spark 3.5.1 + JDK 17",
            "16 GB RAM · 16 CPU · không Docker/cluster",
        ]),
        dict(n=12, title="BLOCKER (đã xử lý)", size=11.0, lines=[
            "Thiếu winutils.exe khi Spark ghi Parquet.",
            "transform.py tự set HADOOP_HOME + PATH vào hadoop/bin.",
        ]),
    ], x=X_RIGHT, w=CARD_W)

    s = slides[3]
    drop(s, "Text 2")
    stack(s, [
        dict(n=3, title="SOURCE", size=10.5, mono=True, lines=[
            "Kaggle — Trending YouTube Videos (113 countries)",
            "kaggle.com/datasets/asaniczka/",
            "trending-youtube-videos-113-countries",
            "License ODC-By · tải 2026-10-20",
        ]),
        dict(n=6, title="SCALE / VERSION", size=11.0, lines=[
            "8.591.418.141 bytes (8,59 GB)",
            "6.094.047 record · 18 cột · 113 quốc gia",
            "snapshot 2023-10-26 → 2026-10-20",
        ]),
        dict(n=9, title="PROVENANCE", size=11.0, lines=[
            "Copy nguyên bản + SHA-256 + manifest",
            "Raw không sửa 1 byte · file lớn không commit",
        ]),
    ], x=X_LEFT, w=CARD_W)
    stack(s, [dict(n=12, title="INGESTION EVIDENCE", size=11.0, mono=True, lines=[
        "$ python src\\ingest_file.py --source <csv> \\",
        "      --dataset youtube_trending",
        "",
        "success | files=1",
        "  bytes=8591418141",
        "  sha256=4310930bd2d1…",
        "  run=run_20261020T143013Z",
        "",
        "$ python src\\verify_raw_records.py",
        "  records=6,094,047",
        "  physical_lines=134,442,959",
        "  multi_line=True",
        "",
        ("Exit code = 0", GREEN, True),
    ])], x=X_RIGHT, w=CARD_W)

    s = slides[4]
    drop(s, "Text 2")
    place(by_name(s, "Shape 3"), width=5.80)
    stack(s, [dict(n=3, title="RAW STORAGE", size=10.5, mono=True, w=5.80, lines=[
        "data/raw/youtube_trending/",
        "  ingest_date=2026-10-20/",
        "    run_20261020T143013Z/",
        "      trending_yt_videos_113_countries.csv",
        "        8.591.418.141 bytes",
        "      manifest.json",
        "",
        "Không sửa 1 byte · size + SHA-256 khớp manifest",
    ])], x=X_LEFT, w=5.80)
    stack(s, [dict(n=6, title="MANIFEST / LOG", size=10.5, mono=True, lines=[
        "dataset     : youtube_trending",
        "source      : …/trending_yt_videos_113_countries.csv",
        "ingested_at : 2026-10-20T14:30:13+00:00",
        "bytes       : 8591418141",
        "sha256      : 4310930bd2d1… (64 ký tự)",
        "destination : data/raw/…/run_…/",
        "records     : 6094047 (RFC4180)",
        "status      : success",
        "amendments  : 1 mục (audit trail)",
        "",
        "logs/env_check.txt   → PASS",
        "logs/ingest.log      → success + duplicate_skip",
    ])], x=X_RIGHT, w=CARD_W)

    s = slides[5]
    drop(s, "Text 2")
    drop(s, "Text 15")
    grid(s, [
        dict(n=3, title="CHẠY INGEST 3 LẦN", size=10.0, mono=True, lines=[
            "lần 1: tạo run + copy + manifest",
            "lần 2–3: hash trùng → SKIP",
            "→ raw giữ nguyên, đúng 1 run",
            "",
            "hash khác → run mới (append)",
            "gate fail → exit 1",
        ]),
        dict(n=6, title="BLOCKER LỚN NHẤT", size=10.5, lines=[
            "Schema nguồn thiếu category_id; description nhiều dòng → Spark đếm sai 17 lần.",
            "Đã xử lý: SRC_MAP + split_csv.py + verify_raw_records.py.",
            "Owner: Phúc · deadline: trước Buổi 7.",
        ]),
        dict(n=9, title="EVIDENCE LINKS", size=10.0, mono=True, lines=[
            "tag: checkpoint-2",
            "manifest: data/raw/…/run_…/manifest.json",
            "logs/env_check.txt",
            "logs/ingest.log",
            "docs/Checkpoint2_Report_*.pptx",
            "docs/Checkpoint2_Brief_Filled.docx",
        ]),
        dict(n=12, title="TRƯỚC BUỔI 7", size=11.0, lines=[
            "1) Ingest snapshot mới để so sánh growth 2 kỳ.",
            "2) Chốt đề cương Đồ án 3.",
        ]),
    ])

    for part in [*prs.slides, prs.slide_master, *prs.slide_master.slide_layouts]:
        for el in part._element.iter():
            if el.get("typeface") == "Aptos":
                el.set("typeface", BODY_FONT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"OK -> {OUT} (tag={tag})")
    return validate(slides)


def validate(slides) -> int:
    problems: list[str] = []
    shapes = {i: {sh.name: sh for sh in s.shapes} for i, s in enumerate(slides, 1)}

    for idx, slide in enumerate(slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            size = max((r.font.size.pt for r in runs if r.font.size), default=18.0)
            mono = any(r.font.name == MONO_FONT for r in runs)
            lines = [p.text for p in shape.text_frame.paragraphs]
            need = measure(lines, shape.width / EMU_IN, size, mono)
            have = shape.height / EMU_IN
            if need > have + 0.03:
                problems.append(f"slide {idx} · {shape.name}: chữ {need:.2f}in > khung {have:.2f}in")

        for a, b in (("Shape 3", "Shape 6"), ("Shape 3", "Shape 9"), ("Shape 3", "Shape 12"),
                     ("Shape 6", "Shape 9"), ("Shape 6", "Shape 12"), ("Shape 9", "Shape 12")):
            if a not in shapes[idx] or b not in shapes[idx]:
                continue
            ra, rb = shapes[idx][a], shapes[idx][b]
            overlap = (
                ra.left < rb.left + rb.width and rb.left < ra.left + ra.width
                and ra.top < rb.top + rb.height and rb.top < ra.top + ra.height
            )
            if overlap:
                problems.append(f"slide {idx}: '{a}' chồng '{b}'")

    full = "\n".join(
        sh.text_frame.text for slide in slides for sh in slide.shapes if sh.has_text_frame
    )
    for needle in ("6.094.047", "8.591.418.141", "4310930bd2d1", "run_20261020T143013Z",
                   "duplicate_skip", "checkpoint-2", "env_check"):
        if needle not in full:
            problems.append(f"thiếu nội dung '{needle}'")
    for slide in slides:
        if "Aptos" in slide._element.xml:
            problems.append("còn font Aptos")
        for ph in ("_____", "DÁN ", "TODO", "________"):
            if ph in "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame):
                problems.append(f"còn placeholder {ph!r}")

    if problems:
        print("\nLAYOUT CÓ VẤN ĐỀ:")
        for p in problems:
            print("  - " + p)
        return 1
    print("LAYOUT OK: không tràn khung, không chồng nhau, đủ nội dung.")
    return 0


if __name__ == "__main__":
    sys.exit(build())
